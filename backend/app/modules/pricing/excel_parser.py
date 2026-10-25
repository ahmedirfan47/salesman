"""Parse an uploaded .xlsx price sheet into validated price lines.

Real-world sheets have inconsistent headers, so we match each column header
against a set of known aliases (English + Arabic). Anything we cannot read
with certainty fails loudly with the row number, rather than being guessed --
protecting the 'never hallucinate prices' rule at the ingestion layer.

Rules:
- A "cost" / "purchase price" column is NEVER used as the selling price. When
  present, it is imported as the product's cost (used for margin analytics).
- If two columns could both be the price (or SKU, or quantity), the upload is
  rejected instead of silently picking one.
- Numbers may use Arabic-Indic digits, thousands separators and a currency
  label ("SAR 1,050", "٣٥ ر.س"). Zero, negative or non-numeric prices are
  rejected.
- Currency comes from a currency column, a label in the price cell, or the
  price header ("Price (SAR)"), in that order. One currency per sheet.
"""
from __future__ import annotations

import io
import re
import unicodedata
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.schemas.pricing import PriceLineIn

# header aliases (normalised) -> canonical field
_ALIASES: dict[str, set[str]] = {
    "sku": {"sku", "code", "item code", "product code", "item no", "item number", "id",
            "الرمز", "الكود", "رمز المنتج", "كود المنتج"},
    "name": {"name", "product", "product name", "item", "item name", "description",
             "الاسم", "المنتج", "اسم المنتج", "الصنف"},
    "price": {"price", "unit price", "selling price", "sell price", "sale price", "rate",
              "السعر", "الثمن", "سعر البيع", "سعر الوحدة"},
    # Never mistaken for the selling price; imported as the product's cost.
    "cost": {"cost", "unit cost", "cost price", "purchase price", "buying price", "buy price",
             "التكلفة", "سعر التكلفة", "سعر الشراء"},
    "quantity": {"quantity", "qty", "stock", "available", "quantity available", "stock qty",
                 "الكمية", "المخزون", "الكمية المتاحة"},
    "currency": {"currency", "ccy", "العملة"},
}

# currency labels found in headers or cells -> ISO 4217 code
_CURRENCY_TOKENS: list[tuple[str, str]] = [
    ("ر.س", "SAR"), ("ريال", "SAR"), ("﷼", "SAR"), ("sar", "SAR"), ("sr", "SAR"),
    ("د.إ", "AED"), ("درهم", "AED"), ("aed", "AED"),
    ("pkr", "PKR"), ("rs.", "PKR"), ("rs", "PKR"),
    ("usd", "USD"), ("us$", "USD"), ("$", "USD"),
]
_ISO = re.compile(r"^[A-Z]{3}$")

_DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹٫٬", "01234567890123456789.,")
_THOUSANDS = re.compile(r"^-?\d{1,3}(,\d{3})+(\.\d+)?$")

_MAX_ERRORS_SHOWN = 20


class ExcelParseError(Exception):
    """Raised when the sheet can't be parsed or a required column is missing."""


def _normalise(text: object) -> str:
    s = unicodedata.normalize("NFKC", str(text)).replace("\u00a0", " ")
    s = re.sub(r"[\u200e\u200f\u202a-\u202e]", "", s)  # bidi marks pasted from Arabic text
    return re.sub(r"\s+", " ", s).strip().lower()


def _find_currency(text: str) -> str | None:
    t = _normalise(text)
    for token, code in _CURRENCY_TOKENS:
        if token.isascii() and token.isalpha():
            if re.search(rf"(?<![a-z]){re.escape(token)}(?![a-z])", t):
                return code
        elif token in t:
            return code
    return None


def _canonical(header: object) -> tuple[str | None, str | None]:
    """Return (field, currency hint) for a header cell, e.g. 'Price (SAR)' -> ('price', 'SAR')."""
    if header is None:
        return None, None
    raw = _normalise(header)
    hint = None
    m = re.search(r"[(\[](.*?)[)\]]", raw)
    if m:
        hint = _find_currency(m.group(1))
        raw = (raw[: m.start()] + raw[m.end():]).strip()
    raw = raw.strip(" :*-_")
    for field, aliases in _ALIASES.items():
        if raw in aliases:
            return field, hint
    return None, None


def _to_decimal(value: object) -> tuple[Decimal | None, str | None]:
    """Parse a numeric cell. Returns (number, currency found in the cell). number is None if unreadable."""
    if value is None or isinstance(value, bool):
        return None, None
    if isinstance(value, (int, float)):
        try:
            d = Decimal(str(value))
        except InvalidOperation:
            return None, None
        return (d if d.is_finite() else None), None

    text = unicodedata.normalize("NFKC", str(value)).translate(_DIGITS).strip()
    ccy = _find_currency(text)
    cleaned = text.lower()
    for token, _ in _CURRENCY_TOKENS:
        cleaned = cleaned.replace(token, " ")
    cleaned = re.sub(r"\s+", "", cleaned)
    if _THOUSANDS.match(cleaned):
        cleaned = cleaned.replace(",", "")
    if not cleaned or "," in cleaned:
        return None, ccy
    try:
        d = Decimal(cleaned)
    except InvalidOperation:
        return None, ccy
    return (d if d.is_finite() else None), ccy


def _to_sku(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel stores 1001 as 1001.0; keep the SKU as "1001"
    s = unicodedata.normalize("NFKC", str(value)).translate(_DIGITS).strip()
    return s or None


def _cell(raw: tuple, idx: int | None):
    if idx is None or idx >= len(raw):
        return None
    return raw[idx]


def parse_price_sheet(file_bytes: bytes) -> list[PriceLineIn]:
    try:
        wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ExcelParseError(f"Could not open the file as an Excel workbook: {exc}") from exc

    try:
        return _parse(wb.active)
    finally:
        wb.close()


def _parse(ws) -> list[PriceLineIn]:
    rows = ws.iter_rows(values_only=True)

    # 1. locate the header row: the first row with a recognisable SKU and price column
    header_row_num = 0
    columns: dict[str, list[tuple[int, str]]] = {}
    price_hint: str | None = None
    cost_hint: str | None = None
    for row_num, raw in enumerate(rows, start=1):
        found: dict[str, list[tuple[int, str]]] = {}
        hints: dict[str, str] = {}
        for idx, cell in enumerate(raw):
            field, hint = _canonical(cell)
            if field:
                found.setdefault(field, []).append((idx, str(cell).strip()))
                if hint:
                    hints[field] = hint
        if "sku" in found and "price" in found:
            header_row_num, columns = row_num, found
            price_hint, cost_hint = hints.get("price"), hints.get("cost")
            break

    if not header_row_num:
        raise ExcelParseError(
            "Could not find a header row. The sheet must have columns for at least "
            "SKU and Price (recognised names include: sku/code, price/unit price/selling price, "
            "or الرمز and السعر)."
        )

    # A generic "ID" column (often just a row number) gives way to a real SKU/code column.
    skus = columns.get("sku", [])
    if len(skus) > 1:
        specific = [c for c in skus if _normalise(c[1]) != "id"]
        if specific:
            columns["sku"] = specific

    for field in ("sku", "price", "cost", "quantity", "currency"):
        if len(columns.get(field, [])) > 1:
            names = ", ".join(f"'{h}' (column {get_column_letter(i + 1)})" for i, h in columns[field])
            raise ExcelParseError(
                f"More than one column could be the {field}: {names}. "
                "Rename or remove the extra column so the sheet is unambiguous."
            )

    if "quantity" not in columns:
        raise ExcelParseError(
            "Missing required column: quantity (recognised names include: quantity, qty, stock, الكمية)."
        )

    col = {f: v[0][0] for f, v in columns.items()}

    # 2. read the product rows
    lines: list[PriceLineIn] = []
    errors: list[str] = []
    seen: dict[str, int] = {}
    currencies: set[str] = set()

    for row_num, raw in enumerate(rows, start=header_row_num + 1):
        sku = _to_sku(_cell(raw, col["sku"]))
        if sku is None:
            continue  # blank row

        where = f"Row {row_num} (SKU {sku})"
        if len(sku) > 64:
            errors.append(f"{where}: SKU is longer than 64 characters")
            continue
        if sku in seen:
            errors.append(f"{where}: SKU appears more than once (first on row {seen[sku]})")
            continue
        seen[sku] = row_num

        price_cell = _cell(raw, col["price"])
        price, cell_ccy = _to_decimal(price_cell)
        if price is None:
            shown = "empty" if price_cell in (None, "") else f"'{price_cell}'"
            errors.append(f"{where}: price is {shown}, not a number")
            continue
        if price <= 0:
            errors.append(f"{where}: price must be greater than 0")
            continue
        if price.as_tuple().exponent < -2:  # e.g. a formula result like 10/3
            price = price.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

        unit_cost, cost_ccy = None, None
        if "cost" in col:
            cost_cell = _cell(raw, col["cost"])
            if cost_cell is not None and str(cost_cell).strip():
                unit_cost, cost_ccy = _to_decimal(cost_cell)
                cost_ccy = cost_ccy or cost_hint
                if unit_cost is None:
                    errors.append(f"{where}: cost '{cost_cell}' is not a number")
                    continue
                if unit_cost < 0:
                    errors.append(f"{where}: cost cannot be negative")
                    continue
                if unit_cost.as_tuple().exponent < -2:
                    unit_cost = unit_cost.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)

        qty_cell = _cell(raw, col["quantity"])
        if qty_cell is None or (isinstance(qty_cell, str) and not qty_cell.strip()):
            quantity = 0
        else:
            qty, _ = _to_decimal(qty_cell)
            if qty is None or qty != qty.to_integral_value():
                errors.append(f"{where}: quantity '{qty_cell}' is not a whole number")
                continue
            quantity = int(qty)
            if quantity < 0:
                errors.append(f"{where}: quantity cannot be negative")
                continue

        currency = None
        if "currency" in col:
            c = _cell(raw, col["currency"])
            if c not in (None, ""):
                currency = _find_currency(str(c)) or str(c).strip().upper()
        currency = currency or cell_ccy or price_hint or cost_ccy
        if currency is not None:
            if not _ISO.match(currency):
                errors.append(f"{where}: currency '{currency}' is not a 3-letter code such as SAR")
                continue
            currencies.add(currency)
        if cost_ccy and cost_ccy != (currency or "USD"):
            errors.append(f"{where}: cost is in {cost_ccy} but the price is in {currency or 'USD'}")
            continue

        name = None
        if "name" in col:
            n = _cell(raw, col["name"])
            name = str(n).strip() if n is not None and str(n).strip() else None

        line = {"sku": sku, "unit_price": price, "quantity": quantity, "name_en": name, "unit_cost": unit_cost}
        if currency:
            line["currency"] = currency
        lines.append(PriceLineIn(**line))

    if len(currencies) > 1:
        errors.append(
            f"The sheet mixes currencies ({', '.join(sorted(currencies))}). Use one currency per price sheet."
        )

    if errors:
        shown = errors[:_MAX_ERRORS_SHOWN]
        more = len(errors) - len(shown)
        tail = f"\n- …and {more} more" if more else ""
        raise ExcelParseError("Problems found in the sheet:\n- " + "\n- ".join(shown) + tail)
    if not lines:
        raise ExcelParseError("No product rows found in the sheet.")

    return lines