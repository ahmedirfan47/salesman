"""Read a prospect list from .xlsx or .csv.

Headers are matched against English and Arabic aliases, so a sales team's own
spreadsheet works without reformatting. Only a name column is required. Rows
that can't be read are reported with their row number; the rest are imported.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata

from openpyxl import load_workbook

from app.modules.prospects.phones import InvalidPhone, normalize_phone

MAX_ROWS = 5000

_ALIASES: dict[str, set[str]] = {
    "name": {"name", "shop", "shop name", "store", "store name", "business", "business name",
             "customer", "customer name", "company", "الاسم", "اسم المحل", "المحل", "اسم المتجر",
             "المتجر", "اسم العميل", "العميل"},
    "phone": {"phone", "phone number", "mobile", "mobile number", "whatsapp", "whatsapp number",
              "number", "contact number", "tel", "telephone", "الجوال", "رقم الجوال", "الهاتف",
              "رقم الهاتف", "واتساب", "رقم الواتساب", "الرقم"},
    "contact_name": {"contact", "contact name", "contact person", "owner", "owner name", "person",
                     "المسؤول", "صاحب المحل", "اسم المسؤول", "المالك"},
    "address": {"address", "location", "street", "العنوان", "الموقع"},
    "area": {"area", "city", "district", "neighborhood", "neighbourhood", "region",
             "المدينة", "الحي", "المنطقة"},
    "zone": {"zone", "delivery zone", "route", "زون", "منطقة التوصيل", "خط التوزيع"},
    "category": {"category", "type", "shop type", "segment", "النوع", "الفئة", "التصنيف", "نوع المحل"},
    "notes": {"notes", "note", "comments", "comment", "remarks", "ملاحظات", "ملاحظة"},
}


class ProspectFileError(Exception):
    """The file can't be read, or it has no recognisable name column."""


def _norm(text: object) -> str:
    s = unicodedata.normalize("NFKC", str(text)).replace("\u00a0", " ")
    s = re.sub(r"[\u200e\u200f\u202a-\u202e]", "", s)  # bidi marks pasted from Arabic text
    return re.sub(r"\s+", " ", s).strip().lower().strip(" :*-_")


def _field(header: object) -> str | None:
    if header is None:
        return None
    h = _norm(header)
    for field, aliases in _ALIASES.items():
        if h in aliases:
            return field
    return None


def _clean(value: object, limit: int) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = re.sub(r"\s+", " ", str(value)).strip()
    return s[:limit] or None


def _rows_from_xlsx(content: bytes) -> list[tuple]:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ProspectFileError(f"Could not open the file as an Excel workbook: {exc}") from exc
    try:
        return [tuple(r) for r in wb.active.iter_rows(values_only=True)]
    finally:
        wb.close()


def _rows_from_csv(content: bytes) -> list[tuple]:
    for encoding in ("utf-8-sig", "cp1256"):  # cp1256: Arabic Windows exports
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ProspectFileError("Could not read the CSV file. Save it as UTF-8 and try again.")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [tuple(r) for r in csv.reader(io.StringIO(text), dialect)]


def parse_prospect_file(content: bytes, filename: str, default_country_code: str) -> tuple[list[dict], list[dict]]:
    """Returns (rows, problems). Each row: {row, name, phone, contact_name, address, area, zone, category, notes}."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        raw_rows = _rows_from_xlsx(content)
    elif name.endswith(".csv"):
        raw_rows = _rows_from_csv(content)
    else:
        raise ProspectFileError("Upload an .xlsx or .csv file.")

    header_idx, col = None, {}
    for i, raw in enumerate(raw_rows[:10]):
        found: dict[str, int] = {}
        for j, cell in enumerate(raw):
            f = _field(cell)
            if f and f not in found:
                found[f] = j
        if "name" in found:
            header_idx, col = i, found
            break
    if header_idx is None:
        raise ProspectFileError(
            "Could not find a header row with a name column "
            "(recognised names include: name, shop name, store, الاسم, اسم المحل)."
        )

    rows: list[dict] = []
    problems: list[dict] = []
    for i, raw in enumerate(raw_rows[header_idx + 1:], start=header_idx + 2):
        def cell(field: str, limit: int) -> str | None:
            j = col.get(field)
            return _clean(raw[j], limit) if j is not None and j < len(raw) else None

        shop = cell("name", 255)
        if not shop:
            continue  # blank row
        if len(rows) >= MAX_ROWS:
            problems.append({"row": i, "name": shop, "reason": f"Only the first {MAX_ROWS} rows are imported"})
            break

        phone = None
        j = col.get("phone")
        raw_phone = raw[j] if j is not None and j < len(raw) else None
        try:
            phone = normalize_phone(raw_phone, default_country_code)
        except InvalidPhone:
            problems.append({"row": i, "name": shop, "reason": f"Phone '{raw_phone}' is not a valid number"})
            continue

        rows.append({
            "row": i,
            "name": shop,
            "phone": phone,
            "contact_name": cell("contact_name", 255),
            "address": cell("address", 500),
            "area": cell("area", 120),
            "zone": cell("zone", 120),
            "category": cell("category", 64),
            "notes": cell("notes", 2000),
        })
    return rows, problems